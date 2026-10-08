# SPDX-License-Identifier: GPL-3.0-or-later
"""Testy eksportu: dane w języku raportu, dane prywatne, każdy format
wczytany z powrotem (CSV, XLSX, XLS, HTML), okno eksportu."""
import csv
import shutil
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path
from unittest import mock

from baza.repozytoria import slowniki
from raporty import dane
from testy.test_gui import EKRAN, TestPrzebudowy
from testy.test_uslugi import Kolekcja
from uslugi import eksport, inwentarz
from wyjatki import BladNornicy


class Eksport(Kolekcja):
    """Kolekcja z testów usług + model z parametrami, cena, historia, znaki specjalne."""

    def setUp(self):
        super().setUp()
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp)
        m = inwentarz.zapisz_model(self.db, {
            "kategoria_id": self.kat("komputery_fabryczne"), "producent": "Atari", "nazwa": "65XE",
            "rok_od": 1985, "rok_do": 1992, "atrybuty": {"procesor": "6502C", "taktowanie_mhz": 1.77},
            "zlacza": {(self.db.execute("SELECT id FROM zlacze_typ WHERE kod='sio_atari'").fetchone()[0],
                        "posiada"): 1}})
        self.atari = self.nowy(model_id=m, lokalizacja_id=self.polka, numer_seryjny="X 006564",
                               cena_minor=25000, waluta_kod="PLN", zrodlo_nabycia="OLX",
                               uwagi='Klawiatura „<b>” & spacja; "cytat"')
        inwentarz.zapisz_zdarzenie(self.db, {"egzemplarz_id": self.atari, "typ": "test", "wynik": "ok",
                                             "data": "2026-09-30", "opis": "Działa"})

    def opcje(self, **zmiany):
        return dane.Opcje(**{"jezyk": "pl", **zmiany})


class TestDanych(Eksport):
    def test_naglowki_i_wartosci_w_jezyku_raportu(self):
        for jezyk, kolumna, status in (("pl", "Kod inwentarzowy", "Działa"), ("en", "Inventory code", "Working")):
            tabela = dane.tabela_egzemplarzy(self.db, [self.atari], self.opcje(jezyk=jezyk))
            self.assertEqual(kolumna, tabela.kolumny[0].naglowek)
            self.assertIn(status, tabela.wiersze[0])

    def test_parametry_i_zlacza_jako_tekst(self):
        wiersz = dane.tabela_egzemplarzy(self.db, [self.atari], self.opcje()).wiersze[0]
        self.assertIn("Procesor: 6502C; Taktowanie procesora: 1,77 MHz", wiersz)
        self.assertIn("Posiada: 1× Atari SIO", wiersz)

    def test_dane_prywatne_tylko_na_zyczenie(self):
        bez = dane.tabela_egzemplarzy(self.db, [self.atari], self.opcje())
        self.assertNotIn("Cena", [k.naglowek for k in bez.kolumny])
        self.assertFalse(any("OLX" == w for w in bez.wiersze[0]))
        z = dane.tabela_egzemplarzy(self.db, [self.atari], self.opcje(prywatne=True))
        self.assertIn(Decimal("250.00"), z.wiersze[0])

    def test_zakres_lokalizacji_z_czesciami(self):
        w_regale = dane.egzemplarze_w_zakresie(self.db, self.regal)
        self.assertEqual({self.pc, self.plyta, self.cpu, self.atari}, set(w_regale))
        modele = dane.tabela_modeli(self.db, w_regale, self.opcje())
        self.assertEqual(["Atari"], [w[0] for w in modele.wiersze])

    def test_lokalizacja_zamontowanej_czesci(self):
        tabela = dane.tabela_egzemplarzy(self.db, [self.cpu], self.opcje())
        kolumna = [k.naglowek for k in tabela.kolumny].index("Lokalizacja")
        self.assertEqual("Regał A › Półka 2", tabela.wiersze[0][kolumna])


class TestFormatow(Eksport):
    def eksportuj(self, formaty, **opcje):
        return eksport.eksportuj(self.db, self.tmp / "kolekcja", formaty, self.opcje(**opcje))

    def test_csv_po_polsku_dla_excela(self):
        pliki = self.eksportuj(["csv"], prywatne=True)
        self.assertEqual(["kolekcja_egzemplarze.csv", "kolekcja_modele.csv", "kolekcja_lokalizacje.csv",
                          "kolekcja_historia.csv"], [p.name for p in pliki])
        surowe = pliki[0].read_bytes()
        self.assertTrue(surowe.startswith(b"\xef\xbb\xbf"))                  # BOM dla Excela
        with open(pliki[0], encoding="utf-8-sig", newline="") as plik:
            wiersze = list(csv.reader(plik, delimiter=";"))
        atari = next(w for w in wiersze if "65XE" in w[3])
        self.assertIn("250,00", atari)                                       # przecinek dziesiętny
        self.assertIn('Klawiatura „<b>” & spacja; "cytat"', atari)          # średnik i cudzysłów w polu

    def test_csv_po_angielsku(self):
        pliki = self.eksportuj(["csv"], jezyk="en", prywatne=True, tabele=("egzemplarze",))
        self.assertEqual("kolekcja_items.csv", pliki[0].name)       # nazwę bazową podaje użytkownik
        with open(pliki[0], encoding="utf-8-sig", newline="") as plik:
            wiersze = list(csv.reader(plik, delimiter=","))
        self.assertIn("250.00", next(w for w in wiersze if "65XE" in w[3]))

    def test_xlsx(self):
        from openpyxl import load_workbook
        (plik,) = self.eksportuj(["xlsx"], prywatne=True)
        skoroszyt = load_workbook(plik)
        self.assertEqual(["Egzemplarze", "Modele", "Lokalizacje", "Testy i naprawy"], skoroszyt.sheetnames)
        arkusz = skoroszyt["Egzemplarze"]
        naglowki = [c.value for c in arkusz[1]]
        wiersz = next(r for r in arkusz.iter_rows(min_row=2, values_only=True) if r[3] == "65XE")
        cena = wiersz[naglowki.index("Cena")]
        self.assertEqual(250.0, cena)                                        # liczba, nie tekst
        komorka = arkusz.cell(row=[r[3] for r in arkusz.iter_rows(values_only=True)].index("65XE") + 1,
                              column=naglowki.index("Cena") + 1)
        self.assertEqual("#,##0.00", komorka.number_format)
        self.assertEqual("A2", arkusz.freeze_panes)

    def test_xls(self):
        import xlrd
        (plik,) = self.eksportuj(["xls"], prywatne=True)
        skoroszyt = xlrd.open_workbook(str(plik))
        self.assertEqual(["Egzemplarze", "Modele", "Lokalizacje", "Testy i naprawy"], skoroszyt.sheet_names())
        arkusz = skoroszyt.sheet_by_name("Egzemplarze")
        naglowki = arkusz.row_values(0)
        wiersz = next(arkusz.row_values(r) for r in range(1, arkusz.nrows) if arkusz.row_values(r)[3] == "65XE")
        self.assertEqual(250.0, wiersz[naglowki.index("Cena")])
        self.assertEqual("Półka 2", wiersz[naglowki.index("Lokalizacja")].split(" › ")[-1])

    def test_xls_limit_sprawdzany_przed_zapisem(self):
        from raporty import eksport_xls
        tabela = dane.Tabela("x", "Duża", "duza", [dane.Kolumna("A")], [["a"]] * 70000)
        with self.assertRaises(BladNornicy) as k:
            eksport_xls.zapisz([tabela], self.tmp / "duza.xls")
        self.assertEqual("blad.xls_limit", k.exception.klucz)
        self.assertFalse((self.tmp / "duza.xls").exists())                  # nic nie zapisano w połowie

    def test_html(self):
        (plik,) = self.eksportuj(["html"], prywatne=True)
        tresc = plik.read_text(encoding="utf-8")
        self.assertIn('<html lang="pl">', tresc)
        self.assertIn("NORNICA — kolekcja", tresc)
        self.assertIn("Klawiatura „&lt;b&gt;” &amp; spacja", tresc)         # znaki HTML zabezpieczone
        self.assertNotIn("<b>”", tresc)
        self.assertIn("Wartość kolekcji", tresc)
        self.assertIn("250,00", tresc)

    def test_html_bez_danych_prywatnych(self):
        (plik,) = self.eksportuj(["html"])
        tresc = plik.read_text(encoding="utf-8")
        self.assertNotIn("Wartość kolekcji", tresc)
        self.assertNotIn("250", tresc)
        self.assertNotIn("OLX", tresc)

    def test_html_z_miniatura(self):
        from PIL import Image
        import os
        from uslugi import pliki
        zdjecie = self.tmp / "a.png"
        Image.new("RGB", (300, 200), (10, 120, 200)).save(zdjecie)
        with mock.patch.dict(os.environ, {"NORNICA_DANE": str(self.tmp / "dane")}):
            z = pliki.dodaj_plik(self.db, zdjecie, {"egzemplarz_id": self.atari})
            pliki.ustaw_zdjecie_glowne(self.db, self.atari, z)
            (plik,) = self.eksportuj(["html"], tabele=("egzemplarze",))
        self.assertEqual(1, plik.read_text(encoding="utf-8").count('src="data:image/png;base64,'))

    def test_wszystkie_formaty_naraz_i_rozszerzenie_w_nazwie(self):
        pliki = eksport.eksportuj(self.db, self.tmp / "zbior.xlsx", ["html", "xls", "xlsx"],
                                  self.opcje(tabele=("egzemplarze",)))
        self.assertEqual(["zbior.xlsx", "zbior.xls", "zbior.html"], [p.name for p in pliki])

    def test_bez_formatu_i_tabel(self):
        with self.assertRaises(BladNornicy):
            self.eksportuj([])
        with self.assertRaises(BladNornicy):
            self.eksportuj(["csv"], tabele=())


@unittest.skipUnless(EKRAN, "brak ekranu (uruchom przez xvfb-run)")
class TestOknaEksportu(TestPrzebudowy):
    def setUp(self):
        super().setUp()
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp)

    def test_eksport_z_okna_po_przebudowie(self):
        self.okno.odswiez(zaznacz=f"L{self.polka}")
        okno = self.okno.eksportuj()
        okno.zakres.set("lokalizacja")
        okno.formaty["csv"].set(True)
        okno.tabele["historia"].set(False)
        okno.jezyk.ustaw("en")                      # angielski raport z polskiego okna
        self.okno.przelacz_jezyk("pl")
        self.okno.update()
        pliki = okno.eksportuj(str(self.tmp / "polka"))
        nazwy = {p.name for p in pliki}
        self.assertIn("polka_items.csv", nazwy)
        self.assertIn("polka.xlsx", nazwy)
        self.assertNotIn("polka_history.csv", nazwy)
        self.assertIn(str(self.tmp), okno._etykieta_wyniku.cget("text"))
        from baza.repozytoria import ustawienia
        self.assertIn("csv", ustawienia.pobierz(self.db, "eksport_formaty"))

    def test_dane_prywatne_zawsze_wylaczone_na_starcie(self):
        okno = self.okno.eksportuj()
        okno.prywatne.set(True)
        okno.eksportuj(str(self.tmp / "a"))
        okno.zamknij()
        self.assertFalse(self.okno.eksportuj().prywatne.get())

    def test_brak_xlwt_wylacza_format(self):
        with mock.patch("raporty.eksport_xls.dostepny", return_value=False):
            okno = self.okno.eksportuj()
        self.assertFalse(okno.formaty["xls"].get())
        self.assertFalse(okno.dostepne["xls"])

    test_stan_przezywa_przelaczenie_tam_i_z_powrotem = None
    test_po_przebudowie_nic_nie_siega_po_zniszczone_widzety = None
    test_otwarte_okno_przezywa_przebudowe_z_danymi = None
    test_dodanie_egzemplarza_przez_okno = None
    test_blad_walidacji_nie_zamyka_okna = None


if __name__ == "__main__":
    unittest.main()
