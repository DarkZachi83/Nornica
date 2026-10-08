# SPDX-License-Identifier: GPL-3.0-or-later
"""Eksport XLSX (Excel 2007+, LibreOffice): jeden plik, arkusz na tabelę.
Nagłówek pogrubiony i zamrożony, filtr na kolumnach, kwoty jako liczby
z formatem walutowym — Excel pokaże je w ustawieniach regionalnych użytkownika."""
from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from raporty.dane import KWOTA, LICZBA, Tabela
from wyjatki import BladNornicy

_ZAKAZANE_W_NAZWIE = str.maketrans({z: " " for z in "[]:*?/\\"})


def nazwa_arkusza(nazwa: str) -> str:
    """Excel: najwyżej 31 znaków, bez []:*?/\\."""
    return nazwa.translate(_ZAKAZANE_W_NAZWIE).strip()[:31] or "Arkusz"


def zapisz(tabele: list[Tabela], sciezka: Path) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise BladNornicy("blad.brak_openpyxl") from None
    skoroszyt = Workbook()
    skoroszyt.remove(skoroszyt.active)
    naglowek_tlo = PatternFill("solid", fgColor="E4E1D8")
    for tabela in tabele:
        arkusz = skoroszyt.create_sheet(nazwa_arkusza(tabela.nazwa))
        arkusz.append([k.naglowek for k in tabela.kolumny])
        for komorka in arkusz[1]:
            komorka.font = Font(bold=True)
            komorka.fill = naglowek_tlo
            komorka.alignment = Alignment(vertical="center")
        for wiersz in tabela.wiersze:
            arkusz.append([float(w) if isinstance(w, Decimal) else w for w in wiersz])
        for i, kolumna in enumerate(tabela.kolumny, 1):
            litera = get_column_letter(i)
            arkusz.column_dimensions[litera].width = kolumna.szerokosc + 2
            if kolumna.typ in (KWOTA, LICZBA):
                for (komorka,) in arkusz.iter_rows(min_row=2, min_col=i, max_col=i):
                    komorka.number_format = "#,##0.00" if kolumna.typ == KWOTA else "0"
        arkusz.freeze_panes = "A2"
        if tabela.wiersze:
            arkusz.auto_filter.ref = arkusz.dimensions
    skoroszyt.save(sciezka)
    return sciezka
